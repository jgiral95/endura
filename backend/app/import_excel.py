"""Importa el Excel histórico de Endura Human («Endura Human.xlsx») a una base vacía.

Uso (dentro del contenedor api):
    python -m app.import_excel /ruta/Endura_Human.xlsx --dry-run   # muestra el resumen y no guarda nada
    python -m app.import_excel /ruta/Endura_Human.xlsx             # importa

Qué hace:
- Marcas y productos desde «Inventario». La existencia inicial de cada producto es la columna EXISTENCIAS
  (que en el Excel ya descuenta ventas y patrocinios). Se reparte en lotes según VENCIMIENTO 1/2/3, suponiendo
  que lo vendido salió primero de lo que vencía antes.
- Ventas («BD Ventas») y patrocinios («Patrocinios») como historial: no vuelven a mover stock. Su coste se
  estima con el coste actual del producto y queda marcado como estimado.
- Pagos desde «Vr pagado»; lo que falte queda como cartera por cobrar.
- Compras («Compras») como historial recibido, sin crear lotes.
- Gastos desde «Contabilidad» (domicilios, transportadoras nacionales, redes, marketing, logística, obligaciones).
  Las compras de producto y los envíos de importación (Deblex, DHL, Letsbox) no se repiten como gasto porque ya
  forman parte de las compras.
- Socios y aportes desde la tabla APORTANTE / DINERO / ESPECIE de «Contabilidad».
- Datos del negocio para la factura desde la hoja «Factura».

Solo se ejecuta sobre una base sin productos: no mezcla ni duplica datos.
"""
import re
import sys
import unicodedata
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
import psycopg
from psycopg.rows import dict_row

from app.main import DATABASE_URL, ensure_schema, refresh_product

BRAND_NAMES = {
    "226ERS": "226ERS", "GU ENERGY": "GU Energy", "SALTSTICK": "SaltStick", "ACTIMAX": "Actimax", "SKRATCH LABS": "Skratch Labs",
    "NEVERSECOND": "Neversecond", "PRECISION FUEL": "Precision Fuel & Hydration", "SIS": "SiS", "MAURTEN": "Maurten",
    "ENERGY ELITE": "Energy Élite", "GOING NUTRICION": "Going Nutrición", "ENERVIT": "Enervit", "IB SPORTS": "IB Sports",
    "SANTAMADRE": "Santa Madre", "AMACX": "Amacx", "BEET IT": "Beet It", "FORCE FACTOR": "Force Factor",
}
KNOWN_REFERRERS = {"JARDIN", "JUANES", "CAMILA", "VALERIA", "SOCIA", "IG", "NATY", "JORGE", "DEISY"}
Q2 = Decimal("0.01")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", " ", s.upper()).strip()


def brand_name(raw):
    n = norm(raw).replace("GRAFAM ", "")
    return BRAND_NAMES.get(n, str(raw).replace("_", " ").strip().title())


def num(v):
    if v is None or isinstance(v, (datetime, date)):
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    try:
        return Decimal(str(v).replace(",", "").strip())
    except Exception:
        return None


def money(v):
    return Decimal(v).quantize(Q2, ROUND_HALF_UP)


def month_end(d):
    """Los vencimientos se imprimen como mes/año: valen hasta el último día de ese mes."""
    if not d:
        return None
    nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return date.fromordinal(nxt.toordinal() - 1)


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def category_for(name):
    n = " " + norm(name) + " "
    def has(*words):
        return any(f" {w} " in n or (w.endswith("*") and f" {w[:-1]}" in n) for w in words)
    if has("GORRO", "TERMO", "PASTILLERO", "BOTELLA", "VISERA", "MEDIAS", "CAMISETA", "BOLSO", "CINTURON", "SHAKER"):
        return "accessory"
    if has("RECOVERY", "REGEN", "RECUPERACION"):
        return "recovery"
    if has("NITRATE", "SHOT", "REMOLACHA", "BEET", "CAFEINA CAPS"):
        return "performance"
    if has("GUMMIE", "GUMMIES", "CHEWS", "BAR", "BARRA", "GOMITAS"):
        return "chews"
    if has("GEL", "GELES", "60ML", "60 ML") or "GO ENERGY" in n or "GO ISOTONIC ENERGY" in n:
        return "gel"
    if has("SAL", "SALT*", "ELECTROLYT*", "CAPS", "CAPSULAS", "FASTCHEWS", "SALTWABLE", "SODIUM", "SODIO", "PH"):
        return "electrolytes"
    if has("DRINK", "DM", "MIX", "HYDRATION", "BOLSA", "ISOTONIC", "CARB", "SACHETS", "TARRO", "PORCIONES", "PORCION", "HIDRATANTE"):
        return "drink_mix"
    return "other"


def build_lots(total, stock, expiries):
    """Reparte la existencia actual en lotes por vencimiento (FEFO: lo vendido salió de lo que vencía antes)."""
    if stock <= 0:
        return []
    total = max(total or 0, stock)
    lots, explicit = [], 0
    for d, q in expiries:
        q = int(q) if q is not None and q > 0 else None
        lots.append([d, q])
        explicit += q or 0
    rem = max(total - explicit, 0)
    blanks = [i for i, l in enumerate(lots) if l[1] is None]
    for i in blanks:
        lots[i][1] = 0
    if blanks:
        lots[blanks[-1]][1] = rem
    elif rem:
        lots.append([None, rem])
    lots.sort(key=lambda l: (l[0] is None, l[0] or date.max))
    sold = total - stock
    for l in lots:
        t = min(sold, l[1])
        l[1] -= t
        sold -= t
    lots = [l for l in lots if l[1] > 0]
    have = sum(l[1] for l in lots)
    if have != stock:  # datos inconsistentes en el Excel: se ajusta el último lote
        if lots:
            lots[-1][1] += stock - have
        else:
            lots = [[None, stock]]
        lots = [l for l in lots if l[1] > 0]
    return lots


class Importer:
    def __init__(self, c, wb):
        self.c, self.wb = c, wb
        self.warnings = []
        self.brands = {}
        self.products = {}      # sku -> dict(id, cost, price, name)
        self.by_name = {}       # nombre normalizado -> sku
        self.stats = OrderedDict()

    def warn(self, msg):
        self.warnings.append(msg)

    def brand(self, raw):
        name = brand_name(raw)
        if name not in self.brands:
            self.brands[name] = self.c.execute(
                "INSERT INTO brands(name) VALUES(%s) ON CONFLICT(name) DO UPDATE SET is_active=TRUE RETURNING id", (name,)).fetchone()["id"]
        return self.brands[name]

    # --------------------------- negocio ---------------------------
    def business(self):
        if "Factura" not in self.wb.sheetnames:
            return
        ws = self.wb["Factura"]
        def text(cell, strip_prefix=""):
            v = ws[cell].value
            return str(v).replace(strip_prefix, "").strip() if v else ""
        self.c.execute("""UPDATE business_settings SET name=%s,owner=%s,tax_id=%s,address=%s,phone=%s,email=%s,invoice_prefix=%s,
            updated_at=NOW() WHERE id=1""", (text("I3").title() or "Endura Human", text("I4"), text("I5", "NIT"), text("I6"),
                                             text("I7", "Celular:"), text("I8", "Correo:"), text("AB6") or "EH"))

    # --------------------------- inventario ---------------------------
    def inventory(self):
        ws = self.wb["Inventario"]
        n_products = n_units = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            sku, name = (r[1] or "").strip() if isinstance(r[1], str) else r[1], r[2]
            if not sku or not name or len(str(name).strip()) <= 2:
                continue
            sku, name = str(sku).strip().upper(), str(name).strip()
            if sku in self.products:
                self.warn(f"Referencia repetida en Inventario: {sku}; se usa la primera")
                continue
            bid = self.brand(r[0] or "Sin marca")
            per_box = int(num(r[4]) or 1) or 1
            total = int(num(r[5]) or 0)
            stock = int(num(r[8]) or 0)
            if stock < 0:
                self.warn(f"{sku} tenía existencias negativas ({stock}); se importa en 0")
                stock = 0
            cost = money(num(r[9]) or 0)
            price = num(r[10])
            pid = self.c.execute("""INSERT INTO products(sku,name,brand_id,category,units_per_box,unit_cost,sale_price,notes)
                VALUES(%s,%s,%s,%s,%s,%s,%s,'') RETURNING id""",
                                 (sku, name, bid, category_for(name), max(per_box, 1), cost, money(price) if price else None)).fetchone()["id"]
            expiries = []
            for di, qi in ((11, 12), (13, 14), (15, 16)):
                d = as_date(r[di]) if di < len(r) else None
                if d:
                    expiries.append((month_end(d), num(r[qi]) if qi < len(r) else None))
            for d, q in build_lots(total, stock, expiries):
                lot = self.c.execute("""INSERT INTO lots(product_id,expiry_date,quantity,initial_quantity,unit_cost,source,note)
                    VALUES(%s,%s,%s,%s,%s,'initial','Saldo importado del Excel') RETURNING id""", (pid, d, q, q, cost)).fetchone()["id"]
                self.c.execute("""INSERT INTO movements(product_id,lot_id,movement_type,quantity_delta,quantity_after,unit_cost,note)
                    VALUES(%s,%s,'initial',%s,0,%s,'Saldo inicial importado del Excel')""", (pid, lot, q, cost))
            refresh_product(self.c, pid)
            self.c.execute("""UPDATE movements m SET quantity_after=x.run FROM (SELECT id,SUM(quantity_delta) OVER (ORDER BY id) run
                FROM movements WHERE product_id=%s) x WHERE m.id=x.id""", (pid,))
            self.products[sku] = {"id": pid, "cost": cost, "price": price, "name": name, "brand_id": bid}
            self.by_name[norm(name)] = sku
            n_products += 1
            n_units += stock
        self.stats["Marcas"] = len(self.brands)
        self.stats["Productos"] = n_products
        self.stats["Unidades en inventario"] = n_units

    def product_for(self, sku, name, brand_raw):
        sku = str(sku or "").strip().upper()
        if sku in self.products:
            return self.products[sku]
        key = norm(name)
        if key in self.by_name:
            return self.products[self.by_name[key]]
        # Producto vendido que ya no está en Inventario: se crea retirado del catálogo para conservar el historial.
        new_sku = sku or f"HIST-{len(self.products) + 1:03d}"
        bid = self.brand(brand_raw or "Sin marca")
        pid = self.c.execute("""INSERT INTO products(sku,name,brand_id,category,is_active,notes) VALUES(%s,%s,%s,%s,FALSE,
            'Creado al importar ventas históricas') RETURNING id""", (new_sku, str(name or new_sku).strip(), bid, category_for(name))).fetchone()["id"]
        self.warn(f"Producto vendido sin ficha en Inventario: {new_sku} {name} (se creó retirado)")
        self.products[new_sku] = {"id": pid, "cost": Decimal(0), "price": None, "name": name, "brand_id": bid}
        self.by_name[key] = new_sku
        return self.products[new_sku]

    # --------------------------- ventas y patrocinios ---------------------------
    @staticmethod
    def _groups(ws):
        groups = OrderedDict()
        for r in ws.iter_rows(min_row=2, values_only=True):
            no = num(r[0])
            if no is None or not r[7]:
                continue
            groups.setdefault(int(no), []).append(r)
        return groups

    def _insert_sale(self, kind, number, rows, header, subtotal, discount, pct, delivery, total, referrer, reason,
                     delivery_status, note):
        sid = self.c.execute("""INSERT INTO sales(kind,number,sale_date,customer_name,referrer,reason,subtotal,discount_percent,discount,
            delivery_fee,total,delivery_status,imported,note) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,TRUE,%s) RETURNING id""",
                             (kind, number, as_date(header[3]) or date.today(), str(header[5] or "").strip(), referrer, reason,
                              subtotal, pct, discount, delivery, total, delivery_status, note)).fetchone()["id"]
        cost_total = Decimal(0)
        for r in rows:
            p = self.product_for(r[8], r[7], r[6])
            q = int(num(r[9]) or 0)
            if q <= 0:
                self.warn(f"{kind} #{number}: línea sin cantidad ({r[7]}), omitida")
                continue
            price = money(num(r[10]) or 0)
            line_cost = money(p["cost"] * q)
            cost_total += line_cost
            self.c.execute("""INSERT INTO sale_items(sale_id,product_id,quantity,unit_price,line_total,unit_cost,line_cost,cost_estimated)
                VALUES(%s,%s,%s,%s,%s,%s,%s,TRUE)""", (sid, p["id"], q, price, money(price * q), p["cost"], line_cost))
        self.c.execute("UPDATE sales SET cost_total=%s WHERE id=%s", (cost_total, sid))
        return sid

    def sales(self):
        ws = self.wb["BD Ventas"]
        n = paid_total = due = 0
        for number, rows in self._groups(ws).items():
            rows.sort(key=lambda r: num(r[1]) or 0)
            h = rows[0]
            subtotal = money(sum((num(r[10]) or 0) * (num(r[9]) or 0) for r in rows))
            pct = num(h[13]) or Decimal(0)
            disc = num(h[15])
            discount = money(disc) if disc is not None else money(subtotal * pct)
            discount = min(discount, subtotal)
            delivery = money(num(h[16]) or 0)
            total = subtotal - discount + delivery
            rec, comment = str(h[23] or "").strip(), str(h[24] or "").strip()
            paid_raw = money(num(h[18]) or 0)
            settled = comment.lower() == "culminado" and norm(h[21]) != "SI"
            if settled and 0 < paid_raw < total:
                # En el Excel la venta quedó cerrada cobrando menos (redondeo o cortesía): la diferencia es descuento.
                gap = total - paid_raw
                if gap <= subtotal - discount:
                    discount += gap
                    total = paid_raw
            notes = []
            referrer = ""
            if rec:
                if norm(rec) in KNOWN_REFERRERS or (rec.isupper() and len(rec) <= 14 and ":" not in rec):
                    referrer = rec.title() if rec.isupper() and len(rec) > 3 else rec
                else:
                    notes.append(rec)
            low = comment.lower()
            delivery_status = "pending" if "entreg" in low and "culminado" not in low else "delivered"
            if comment and comment.lower() != "culminado":
                notes.append(comment)
            pct_eff = (discount / subtotal * 100).quantize(Decimal("0.001")) if subtotal else Decimal(0)
            sid = self._insert_sale("sale", number, rows, h, subtotal, discount, pct_eff, delivery,
                                    total, referrer, None, delivery_status, " · ".join(notes))
            paid = num(h[18]) or Decimal(0)
            if paid > 0:
                paid = money(paid)
                if paid > total or total - paid < 1:
                    paid = total
                method = {"BANCOLOMBIA": "bancolombia", "EFECTIVO": "cash", "NEQUI": "nequi"}.get(norm(h[20]), "other")
                if paid > 0:
                    self.c.execute("""INSERT INTO sale_payments(sale_id,payment_date,amount,method,note) VALUES(%s,%s,%s,%s,
                        'Importado del Excel')""", (sid, as_date(h[19]) or as_date(h[3]) or date.today(), paid, method))
                paid_total += paid
            if total - paid > 0:
                due += 1
            n += 1
        self.stats["Ventas"] = n
        self.stats["Ventas con saldo pendiente"] = due
        self.stats["Pagos importados"] = money(paid_total)

    def giveaways(self):
        if "Patrocinios" not in self.wb.sheetnames:
            return
        ws = self.wb["Patrocinios"]
        n = 0
        for number, rows in self._groups(ws).items():
            rows.sort(key=lambda r: num(r[1]) or 0)
            h = rows[0]
            who = norm(h[5])
            hint = norm(f"{h[20] or ''} {h[24] or ''}")
            if "DANAD" in who or "DANO" in hint or "VENCID" in hint:
                reason = "damage"
            elif "INVENTARIO" in who or "REDES" in who:
                reason = "internal"
            elif "INTERCAMBIO" in hint or "ALIANZA" in hint:
                reason = "alliance"
            elif any(w in who for w in ("ACTIVACION", "RUN", "CENTER", "CARRERA", "EVENTO", "FIT")):
                reason = "event"
            else:
                reason = "ambassador"
            subtotal = money(sum((num(r[10]) or 0) * (num(r[9]) or 0) for r in rows))
            note = " · ".join(x for x in (str(h[20] or "").strip(), str(h[24] or "").strip()) if x)
            self._insert_sale("giveaway", number, rows, h, subtotal, subtotal, Decimal(100), Decimal(0), Decimal(0), "", reason,
                              "delivered", note)
            n += 1
        self.stats["Patrocinios y salidas"] = n

    # --------------------------- compras ---------------------------
    def purchases(self):
        ws = self.wb["Compras"]
        n, total_all = 0, Decimal(0)
        for r in ws.iter_rows(min_row=2, values_only=True):
            d, store, pesos = as_date(r[0]), r[2], num(r[5])
            if not d or not store or pesos is None or str(store).upper() == "TOTAL":
                continue
            usd = num(r[3])
            fx = money(pesos / usd) if usd else Decimal(1)
            ship = money(num(r[11]) or 0) if len(r) > 11 else Decimal(0)
            other = money(num(r[12]) or 0) if len(r) > 12 else Decimal(0)
            goods = money(pesos)
            tracking = str(r[8] or "").strip()
            self.c.execute("""INSERT INTO purchases(purchase_date,supplier,funded_by,origin,currency,fx_rate,invoice_no,tracking,carrier,
                shipping_cost,other_costs,goods_total,total_cop,status,imported,note,received_at)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'received',TRUE,%s,%s)""",
                           (d, str(store).strip(), str(r[1] or "").strip().title(), str(r[6] or "").strip(), "USD" if usd else "COP", fx,
                            str(r[7] or "").strip(), tracking, str(r[9] or "").strip(), ship, other, goods, goods + ship + other,
                            str(r[13] or "").strip() if len(r) > 13 else "", datetime.combine(d, datetime.min.time())))
            n += 1
            total_all += goods + ship + other
        self.stats["Compras (historial)"] = n
        self.stats["Valor compras"] = money(total_all)

    # --------------------------- contabilidad ---------------------------
    def accounting(self):
        if "Contabilidad" not in self.wb.sheetnames:
            return
        ws = self.wb["Contabilidad"]
        rows = list(ws.iter_rows(values_only=True))
        year = next((int(v) for v in rows[0] if isinstance(v, (int, float))), date.today().year)
        months = {"JUNIO": 6, "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
                  "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5}
        header = rows[1]
        blocks = [(i, months[norm(v)]) for i, v in enumerate(header) if v and norm(v) in months]
        section, n, total = None, 0, Decimal(0)
        skip_sections = {"CAPITAL SOCIOS", "COMPRA PRODUCTO"}
        import_carriers = {"DEBLEX", "DHL", "LETSBOX IT"}
        for r in rows[3:]:
            name = r[0]
            if not name:
                continue
            key = norm(name)
            if key.startswith("TOTAL"):
                break
            if all(v is None for v in r[1:]) and str(name).strip().isupper():
                section = key
                continue
            if section in skip_sections or key in ("INGRESO POR VENTAS", "PENDIENTE DE PAGO"):
                continue
            if section == "VENTAS":
                category = "delivery" if "DOMICILIO" in key else None
            elif section == "TRANSPORTADORAS":
                category = None if key in import_carriers else "shipping"
            elif section and "COMUNICACIONES" in section:
                category = "comms"
            elif section == "MARKETING":
                category = "marketing"
            elif section and "LOGISTICA" in section:
                category = "logistics"
            elif section == "OBLIGACIONES":
                category = "platforms" if "SHOPIFY" in key else "obligations"
            else:
                category = "other"
            if not category:
                continue
            for col, month in blocks:
                amount = num(r[col + 1]) if col + 1 < len(r) else None
                if not amount or amount <= 0:
                    continue
                d = as_date(r[col + 2]) if col + 2 < len(r) else None
                d = d or date(year, month, 1)
                paid_by = ""
                m = re.match(r"Domicilios\s+(.+)", str(name).strip(), re.I)
                if m:
                    paid_by = m.group(1).strip()
                self.c.execute("""INSERT INTO expenses(expense_date,category,description,amount,paid_by,imported,note)
                    VALUES(%s,%s,%s,%s,%s,TRUE,'Importado de Contabilidad')""", (d, category, str(name).strip(), money(amount), paid_by))
                n += 1
                total += money(amount)
        self.stats["Gastos"] = n
        self.stats["Valor gastos"] = total
        # Tabla de aportes: APORTANTE | DINERO | ESPECIE
        pos = None
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                if isinstance(v, str) and norm(v) == "APORTANTE":
                    pos = (i, j)
        if not pos:
            return
        i0, j0 = pos
        n_partners = 0
        for r in rows[i0 + 2:]:
            name = r[j0] if j0 < len(r) else None
            if not name or norm(name).startswith("TOTAL"):
                break
            cash, kind_ = num(r[j0 + 1]) or 0, num(r[j0 + 2]) or 0
            label = str(name).strip().title()
            pid = self.c.execute("INSERT INTO partners(name) VALUES(%s) ON CONFLICT(name) DO UPDATE SET is_active=TRUE RETURNING id",
                                 (label,)).fetchone()["id"]
            for k, amount in (("cash", cash), ("in_kind", kind_)):
                if amount > 0:
                    self.c.execute("""INSERT INTO capital_movements(partner_id,movement_date,kind,amount,note)
                        VALUES(%s,%s,%s,%s,'Aporte acordado según Excel')""", (pid, date(year, 6, 1), k, money(amount)))
            n_partners += 1
        self.stats["Socios"] = n_partners

    def run(self):
        self.business()
        self.inventory()
        self.sales()
        self.giveaways()
        self.purchases()
        self.accounting()


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    if len(args) != 1:
        print("Uso: python -m app.import_excel ARCHIVO.xlsx [--dry-run]")
        raise SystemExit(2)
    wb = openpyxl.load_workbook(args[0], data_only=True)
    ensure_schema()
    with psycopg.connect(DATABASE_URL, row_factory=dict_row) as c:
        n = c.execute("SELECT COUNT(*) n FROM products").fetchone()["n"]
        if n:
            print(f"La base ya tiene {n} productos. El importador solo trabaja sobre una base vacía para no duplicar datos.")
            raise SystemExit(1)
        imp = Importer(c, wb)
        imp.run()
        print("Resumen de la importación" + (" (simulación: no se guardó nada)" if dry else ""))
        for k, v in imp.stats.items():
            print(f"  {k}: {v}")
        if imp.warnings:
            print(f"\nAvisos ({len(imp.warnings)}):")
            for w in imp.warnings:
                print("  - " + w)
        if dry:
            c.rollback()
        else:
            c.commit()
            print("\nImportación guardada.")


if __name__ == "__main__":
    main()
